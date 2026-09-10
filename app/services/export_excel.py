"""Export hasil satu sesi ke file Excel memakai openpyxl."""

import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from ..models import (
    MOD_APPROVED,
    MODE_QUIZ,
    TIPE_MC,
    TIPE_RATING,
    TIPE_WORD_CLOUD,
    Jawaban,
    Partisipan,
    Pertanyaan,
    Sesi,
)
from ..utils import ke_utc, normalisasi_kata

# Warna mengikuti palet aplikasi (lime + ink).
_ISI_JUDUL = PatternFill("solid", fgColor="1A1A17")
_ISI_HEADER = PatternFill("solid", fgColor="C6F135")
_ISI_ZEBRA = PatternFill("solid", fgColor="FAF8F3")
_GARIS = Border(*(Side(style="thin", color="E4DFD2"),) * 4)

_NAMA_TIPE = {TIPE_MC: "Multiple Choice", TIPE_WORD_CLOUD: "Word Cloud", TIPE_RATING: "Rating Scale"}


def _tulis_header(ws, baris: int, kolom: list[str]) -> None:
    for i, judul in enumerate(kolom, start=1):
        sel = ws.cell(row=baris, column=i, value=judul)
        sel.font = Font(bold=True, size=10, color="1A1A17")
        sel.fill = _ISI_HEADER
        sel.border = _GARIS
        sel.alignment = Alignment(vertical="center", wrap_text=True)


def _atur_lebar(ws, lebar: list[int]) -> None:
    for i, w in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _judul_lembar(ws, teks: str, kolom_akhir: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=kolom_akhir)
    sel = ws.cell(row=1, column=1, value=teks)
    sel.font = Font(bold=True, size=13, color="FFFFFF")
    sel.fill = _ISI_JUDUL
    sel.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 28


async def buat_excel(db: AsyncSession, sesi: Sesi) -> bytes:
    """Rangkai workbook berisi ringkasan, hasil per pertanyaan, dan jawaban mentah."""
    pertanyaan = (
        (
            await db.execute(
                select(Pertanyaan)
                .where(Pertanyaan.session_id == sesi.id)
                .options(selectinload(Pertanyaan.daftar_opsi))
                .order_by(Pertanyaan.urutan, Pertanyaan.id)
            )
        )
        .scalars()
        .all()
    )
    peserta = (
        (await db.execute(select(Partisipan).where(Partisipan.session_id == sesi.id))).scalars().all()
    )
    peta_peserta = {p.id: p for p in peserta}
    id_pertanyaan = [q.id for q in pertanyaan]
    jawaban = []
    if id_pertanyaan:
        jawaban = (
            (
                await db.execute(
                    select(Jawaban)
                    .where(Jawaban.question_id.in_(id_pertanyaan))
                    .order_by(Jawaban.question_id, Jawaban.submitted_at)
                )
            )
            .scalars()
            .all()
        )
    per_pertanyaan: dict[int, list[Jawaban]] = {qid: [] for qid in id_pertanyaan}
    for a in jawaban:
        per_pertanyaan.setdefault(a.question_id, []).append(a)

    wb = Workbook()

    # --- Lembar 1: Ringkasan ---------------------------------------------
    ws = wb.active
    ws.title = "Ringkasan"
    _atur_lebar(ws, [26, 60])
    _judul_lembar(ws, f"Ringkasan Sesi — {sesi.judul}", 2)
    baris = [
        ("Kode Sesi", sesi.kode_sesi),
        ("Judul", sesi.judul),
        ("Mode", "Quiz Mode" if sesi.mode == MODE_QUIZ else "Survey Mode"),
        ("Status", sesi.status),
        ("Dibuat", (ke_utc(sesi.created_at) or "").strftime("%d-%m-%Y %H:%M UTC") if sesi.created_at else "-"),
        ("Selesai", (ke_utc(sesi.ended_at) or "").strftime("%d-%m-%Y %H:%M UTC") if sesi.ended_at else "-"),
        ("Jumlah Pertanyaan", len(pertanyaan)),
        ("Jumlah Partisipan", len(peserta)),
        ("Total Jawaban Masuk", len(jawaban)),
    ]
    for i, (label, nilai) in enumerate(baris, start=3):
        sel_label = ws.cell(row=i, column=1, value=label)
        sel_label.font = Font(bold=True, size=10)
        sel_label.fill = _ISI_ZEBRA
        sel_label.border = _GARIS
        sel_nilai = ws.cell(row=i, column=2, value=nilai)
        sel_nilai.border = _GARIS

    # --- Lembar 2: Hasil per pertanyaan ----------------------------------
    ws = wb.create_sheet("Hasil per Pertanyaan")
    _atur_lebar(ws, [6, 46, 18, 34, 12, 12])
    _judul_lembar(ws, "Rekap Hasil per Pertanyaan", 6)
    kursor = 3
    for nomor, q in enumerate(pertanyaan, start=1):
        daftar = per_pertanyaan.get(q.id, [])
        sel = ws.cell(row=kursor, column=1, value=f"{nomor}.")
        sel.font = Font(bold=True)
        sel_q = ws.cell(row=kursor, column=2, value=q.teks)
        sel_q.font = Font(bold=True, size=11)
        ws.cell(row=kursor, column=3, value=_NAMA_TIPE.get(q.tipe, q.tipe)).font = Font(size=9, color="6B6B60")
        kursor += 1

        _tulis_header(ws, kursor, ["", "", "", "Jawaban", "Jumlah", "Persen"])
        kursor += 1

        if q.tipe == TIPE_MC:
            total = len(daftar)
            hitung = Counter(a.option_id for a in daftar if a.option_id is not None)
            for o in q.daftar_opsi:
                n = hitung.get(o.id, 0)
                label = o.teks + ("  (jawaban benar)" if o.is_benar and sesi.mode == MODE_QUIZ else "")
                ws.cell(row=kursor, column=4, value=label).border = _GARIS
                ws.cell(row=kursor, column=5, value=n).border = _GARIS
                ws.cell(row=kursor, column=6, value=round(n * 100 / total, 1) if total else 0).border = _GARIS
                kursor += 1
        elif q.tipe == TIPE_WORD_CLOUD:
            disetujui = [a for a in daftar if a.status_moderasi == MOD_APPROVED and a.teks]
            hitung = Counter(normalisasi_kata(a.teks) for a in disetujui)
            total = sum(hitung.values())
            for kata, n in hitung.most_common():
                ws.cell(row=kursor, column=4, value=kata).border = _GARIS
                ws.cell(row=kursor, column=5, value=n).border = _GARIS
                ws.cell(row=kursor, column=6, value=round(n * 100 / total, 1) if total else 0).border = _GARIS
                kursor += 1
            ditolak = len(daftar) - len(disetujui)
            if ditolak:
                ws.cell(row=kursor, column=4, value=f"(+{ditolak} jawaban belum/tidak disetujui)").font = Font(
                    italic=True, size=9, color="6B6B60"
                )
                kursor += 1
        elif q.tipe == TIPE_RATING:
            nilai = [a.nilai_rating for a in daftar if a.nilai_rating is not None]
            hitung = Counter(nilai)
            total = len(nilai)
            for n in range(1, (q.rating_maks or 5) + 1):
                c = hitung.get(n, 0)
                ws.cell(row=kursor, column=4, value=f"Nilai {n}").border = _GARIS
                ws.cell(row=kursor, column=5, value=c).border = _GARIS
                ws.cell(row=kursor, column=6, value=round(c * 100 / total, 1) if total else 0).border = _GARIS
                kursor += 1
            rata = round(sum(nilai) / total, 2) if total else 0
            sel = ws.cell(row=kursor, column=4, value="Rata-rata")
            sel.font = Font(bold=True)
            ws.cell(row=kursor, column=5, value=rata).font = Font(bold=True)
            kursor += 1

        kursor += 1

    # --- Lembar 3: Jawaban mentah ----------------------------------------
    ws = wb.create_sheet("Jawaban Mentah")
    kolom = [
        "No",
        "Pertanyaan",
        "Tipe",
        "Partisipan",
        "Jawaban",
        "Benar",
        "Waktu Jawab (ms)",
        "Poin",
        "Status Moderasi",
        "Waktu Submit",
    ]
    _atur_lebar(ws, [5, 40, 16, 20, 30, 8, 16, 8, 16, 20])
    _judul_lembar(ws, "Seluruh Jawaban Masuk", len(kolom))
    _tulis_header(ws, 3, kolom)
    baris_ke = 4
    peta_opsi = {o.id: o.teks for q in pertanyaan for o in q.daftar_opsi}
    peta_q = {q.id: q for q in pertanyaan}
    for i, a in enumerate(jawaban, start=1):
        q = peta_q.get(a.question_id)
        p = peta_peserta.get(a.participant_id)
        if a.option_id is not None:
            isi = peta_opsi.get(a.option_id, str(a.option_id))
        elif a.nilai_rating is not None:
            isi = a.nilai_rating
        else:
            isi = a.teks or ""
        nilai_baris = [
            i,
            q.teks if q else "-",
            _NAMA_TIPE.get(q.tipe, q.tipe) if q else "-",
            (p.nickname if p and p.nickname else "Anonim"),
            isi,
            "Ya" if a.benar else "",
            a.waktu_jawab_ms,
            a.poin,
            a.status_moderasi,
            (ke_utc(a.submitted_at) or "").strftime("%d-%m-%Y %H:%M:%S") if a.submitted_at else "",
        ]
        for kol, nilai in enumerate(nilai_baris, start=1):
            sel = ws.cell(row=baris_ke, column=kol, value=nilai)
            sel.border = _GARIS
            if i % 2 == 0:
                sel.fill = _ISI_ZEBRA
        baris_ke += 1
    ws.freeze_panes = "A4"

    # --- Lembar 4: Leaderboard (khusus Quiz Mode) ------------------------
    if sesi.mode == MODE_QUIZ:
        ws = wb.create_sheet("Leaderboard")
        _atur_lebar(ws, [10, 30, 14, 16])
        _judul_lembar(ws, "Leaderboard Final", 4)
        _tulis_header(ws, 3, ["Peringkat", "Nickname", "Total Poin", "Jawaban Benar"])
        benar_per_peserta = Counter(a.participant_id for a in jawaban if a.benar)
        urut = sorted(peserta, key=lambda p: (-(p.total_poin or 0), (p.nickname or "").lower()))
        for i, p in enumerate(urut, start=1):
            nilai_baris = [i, p.nickname or "Anonim", p.total_poin or 0, benar_per_peserta.get(p.id, 0)]
            for kol, nilai in enumerate(nilai_baris, start=1):
                sel = ws.cell(row=3 + i, column=kol, value=nilai)
                sel.border = _GARIS
                if i <= 3:
                    sel.font = Font(bold=True)
                elif i % 2 == 0:
                    sel.fill = _ISI_ZEBRA
        ws.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
